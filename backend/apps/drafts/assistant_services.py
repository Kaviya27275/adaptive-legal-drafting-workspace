import json
from io import BytesIO

import requests
from django.conf import settings
from django.utils import timezone

from apps.compliance.models import DocumentType


DEFAULT_TEMPLATE_BY_TYPE = {
    "Non Disclosure Agreement": """NON-DISCLOSURE AGREEMENT
Date: {today_date}
Jurisdiction: {jurisdiction}

This Non-Disclosure Agreement is entered into on {today_date} between:
1. Disclosing Party: [Insert Name]
2. Receiving Party: [Insert Name]

1. Definitions
Confidential Information means all non-public information disclosed in any form.

2. Obligations
The Receiving Party shall keep all Confidential Information confidential and use it only for the agreed purpose.

3. Exclusions
Information that is public or independently developed is excluded.

4. Term and Termination
This Agreement remains effective for [Insert Period], unless terminated earlier.

5. Governing Law and Jurisdiction
This Agreement is governed by the laws of {jurisdiction}.

Signatures
Disclosing Party: ____________________
Receiving Party: ____________________
""",
    "Employment Contract": """EMPLOYMENT CONTRACT
Date: {today_date}
Jurisdiction: {jurisdiction}

This Employment Contract is made on {today_date} between:
Employer: [Insert Company Name]
Employee: [Insert Employee Name]

1. Position and Duties
Employee is appointed as [Insert Role].

2. Compensation
Monthly salary: [Insert Amount].

3. Working Hours and Leave
Working hours and leave policy shall follow company policy and applicable law.

4. Confidentiality
Employee must maintain confidentiality of employer information.

5. Termination
Either party may terminate with [Insert Notice Period].

6. Governing Law
This Agreement is governed by the laws of {jurisdiction}.

Signatures
Employer: ____________________
Employee: ____________________
""",
    "Service Agreement": """SERVICE AGREEMENT
Date: {today_date}
Jurisdiction: {jurisdiction}

This Service Agreement is made on {today_date} between:
Client: [Insert Name]
Service Provider: [Insert Name]

1. Scope of Services
Service Provider shall provide: [Insert Scope].

2. Fees and Payment
Client shall pay [Insert Amount] as per agreed milestones.

3. Term and Termination
This Agreement begins on {today_date} and continues until completion or termination.

4. Liability and Indemnity
Parties agree to liability and indemnity terms as per applicable law.

5. Governing Law
This Agreement is governed by the laws of {jurisdiction}.

Signatures
Client: ____________________
Service Provider: ____________________
""",
    "Lease Agreement": """LEASE AGREEMENT
Date: {today_date}
Jurisdiction: {jurisdiction}

This Lease Agreement is made on {today_date} between:
Lessor: [Insert Name]
Lessee: [Insert Name]

1. Premises
Lessor leases premises located at [Insert Address].

2. Lease Term
Lease starts on {today_date} for [Insert Term].

3. Rent and Deposit
Monthly rent: [Insert Amount]; Security deposit: [Insert Amount].

4. Maintenance
Maintenance responsibilities are allocated as mutually agreed.

5. Termination
Termination terms and notice period: [Insert Details].

6. Governing Law
This Agreement is governed by the laws of {jurisdiction}.

Signatures
Lessor: ____________________
Lessee: ____________________
""",
    "Legal Notice": """LEGAL NOTICE
Date: {today_date}
Jurisdiction: {jurisdiction}

To:
[Recipient Name]
[Recipient Address]

Subject: Legal Notice regarding [Insert Subject]

Dear Sir/Madam,

Under instructions from and on behalf of my client [Insert Client Name], this legal notice is issued as follows:

1. Facts
[Insert relevant factual background].

2. Breach/Default
[Insert details of breach, default, or grievance].

3. Demand
You are hereby called upon to [perform/payment/cease action] within [Insert Number] days from receipt of this notice.

4. Consequence of Non-Compliance
Failing compliance, my client shall be constrained to initiate appropriate civil/criminal proceedings at your risk as to costs and consequences.

Sincerely,
[Advocate Name]
[Law Firm/Address]
""",
    "Consultancy Agreement": """CONSULTANCY AGREEMENT
Date: {today_date}
Jurisdiction: {jurisdiction}

This Consultancy Agreement is made on {today_date} between:
Client: [Insert Name]
Consultant: [Insert Name]

1. Scope of Services
Consultant shall provide: [Insert Scope].

2. Fees and Payment
Client shall pay [Insert Amount] as per agreed milestones.

3. Term and Termination
This Agreement begins on {today_date} and continues until completion or termination.

4. Confidentiality
Consultant shall keep all Client information confidential.

5. Governing Law
This Agreement is governed by the laws of {jurisdiction}.

Signatures
Client: ____________________
Consultant: ____________________
""",
}

TEMPLATE_NAME_ALIASES = {
    "NDA": "Non Disclosure Agreement",
    "Non-Disclosure Agreement": "Non Disclosure Agreement",
    "Employment Agreement": "Employment Contract",
    "Commercial Lease Agreement": "Lease Agreement",
}


def _extract_txt(uploaded_file) -> str:
    raw = uploaded_file.read()
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("latin-1", errors="ignore")


def _extract_docx(uploaded_file) -> str:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise ValueError("DOCX parsing support is not installed.") from exc
    document = DocxDocument(BytesIO(uploaded_file.read()))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)


def _extract_doc(uploaded_file) -> str:
    try:
        import textract
    except ImportError as exc:
        raise ValueError("DOC parsing support is not installed.") from exc
    content = textract.process(uploaded_file)
    return content.decode("utf-8", errors="ignore")


def _extract_pdf(uploaded_file) -> str:
    try:
        import pdfplumber
    except ImportError as exc:
        raise ValueError("PDF parsing support is not installed.") from exc
    text_parts = []
    with pdfplumber.open(BytesIO(uploaded_file.read())) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_text_from_uploaded_file(uploaded_file) -> str:
    extension = uploaded_file.name.lower().split(".")[-1]
    if extension == "txt":
        return _extract_txt(uploaded_file)
    if extension == "doc":
        return _extract_doc(uploaded_file)
    if extension == "docx":
        return _extract_docx(uploaded_file)
    if extension == "pdf":
        return _extract_pdf(uploaded_file)
    raise ValueError("Unsupported file type. Use .txt, .doc, .docx, or .pdf")


def _fallback_document_type(text: str, allowed_types: list[str]) -> str:
    text_lower = (text or "").lower()
    keyword_map = {
        "nda": "Non Disclosure Agreement",
        "confidential": "Non Disclosure Agreement",
        "employment": "Employment Contract",
        "employee": "Employment Contract",
        "service": "Service Agreement",
        "lease": "Lease Agreement",
        "rent": "Lease Agreement",
        "notice": "Legal Notice",
        "demand": "Legal Notice",
        "breach": "Legal Notice",
    }
    for keyword, mapped_type in keyword_map.items():
        if keyword in text_lower and mapped_type in allowed_types:
            return mapped_type
    return allowed_types[0] if allowed_types else "Unknown"


def detect_document_type_with_ai(text: str) -> dict:
    allowed_types = list(DocumentType.objects.values_list("name", flat=True))
    if not allowed_types:
        allowed_types = list(DEFAULT_TEMPLATE_BY_TYPE.keys())

    prompt = json.dumps(
        {
            "task": "Classify legal document type using content",
            "allowed_document_types": allowed_types,
            "output_schema": {
                "document_type": "string",
                "confidence": 0,
                "reason": "string",
            },
            "document_text": (text or "")[:7000],
            "response_rule": "Return JSON only.",
        }
    )

    try:
        response = requests.post(
            f"{getattr(settings, 'OLLAMA_URL', 'http://localhost:11434')}/api/generate",
            json={
                "model": getattr(settings, "OLLAMA_MODEL", "llama3"),
                "prompt": prompt,
                "stream": False,
                "format": "json",
            },
            timeout=90,
        )
        response.raise_for_status()
        parsed = json.loads(response.json().get("response", "{}"))
        predicted = parsed.get("document_type")
        if predicted not in allowed_types:
            predicted = _fallback_document_type(text, allowed_types)
        return {
            "document_type": predicted,
            "confidence": parsed.get("confidence", 0),
            "reason": parsed.get("reason", ""),
        }
    except (requests.RequestException, ValueError, json.JSONDecodeError):
        return {
            "document_type": _fallback_document_type(text, allowed_types),
            "confidence": 0,
            "reason": "Fallback classifier used because AI detection was unavailable.",
        }


def list_document_templates() -> list[dict]:
    templates = []
    for template_name in DEFAULT_TEMPLATE_BY_TYPE.keys():
        document_type, _ = DocumentType.objects.get_or_create(name=template_name)
        templates.append(
            {
                "document_type_id": document_type.id,
                "document_type_name": template_name,
                "has_template": True,
            }
        )
    return templates


def render_document_template(document_type_name: str, jurisdiction: str = "India") -> str:
    document_type_name = TEMPLATE_NAME_ALIASES.get(document_type_name, document_type_name)
    template = DEFAULT_TEMPLATE_BY_TYPE.get(document_type_name)
    if not template:
        raise ValueError("No template configured for this document type.")
    today = timezone.localdate().isoformat()
    return template.format(today_date=today, jurisdiction=jurisdiction or "India")


def generate_draft_from_template(
    document_type_name: str,
    key_terms: list[str] | None,
    jurisdiction: str = "India",
) -> str:
    template = render_document_template(document_type_name, jurisdiction=jurisdiction)
    terms = [term for term in (key_terms or []) if isinstance(term, str) and term.strip()]
    prompt = json.dumps(
        {
            "task": "Draft a complete legal document using the provided template and key terms.",
            "constraints": [
                "Preserve the overall template structure and headings.",
                "Fill in placeholders and expand sections using key terms.",
                "Return the full draft as plain text.",
            ],
            "document_type": document_type_name,
            "jurisdiction": jurisdiction,
            "key_terms": terms,
            "template": template,
            "response_rule": "Return text only.",
        }
    )

    try:
        response = requests.post(
            f"{getattr(settings, 'OLLAMA_URL', 'http://localhost:11434')}/api/generate",
            json={
                "model": getattr(settings, "OLLAMA_MODEL", "llama3"),
                "prompt": prompt,
                "stream": False,
            },
            timeout=120,
        )
        response.raise_for_status()
        ai_text = response.json().get("response", "") or ""
        return ai_text.strip() or template
    except (requests.RequestException, ValueError, json.JSONDecodeError):
        return template


def drafting_ai_assist(text: str, instruction: str) -> dict:
    prompt = json.dumps(
        {
            "task": "Provide drafting assistance without rewriting full document.",
            "constraints": [
                "No automatic modification of original draft.",
                "Return advisory recommendations only.",
                "Return JSON only.",
            ],
            "output_schema": {
                "clarity_improvements": [],
                "missing_clauses": [],
                "risk_flags": [],
                "recommended_edits": [],
            },
            "instruction": instruction,
            "draft_text": (text or "")[:9000],
        }
    )

    try:
        response = requests.post(
            f"{getattr(settings, 'OLLAMA_URL', 'http://localhost:11434')}/api/generate",
            json={
                "model": getattr(settings, "OLLAMA_MODEL", "llama3"),
                "prompt": prompt,
                "stream": False,
                "format": "json",
            },
            timeout=120,
        )
        response.raise_for_status()
        result = json.loads(response.json().get("response", "{}"))
    except (requests.RequestException, ValueError, json.JSONDecodeError):
        result = {}

    return {
        "clarity_improvements": result.get("clarity_improvements", []),
        "missing_clauses": result.get("missing_clauses", []),
        "risk_flags": result.get("risk_flags", []),
        "recommended_edits": result.get("recommended_edits", []),
    }


def analyze_legal_document(text: str, document_type: str) -> dict:
    prompt = json.dumps(
        {
            "task": "Analyze legal document and return structured insights.",
            "document_type": document_type,
            "output_schema": {
                "detected_document_type": "string",
                "executive_summary": "string",
                "overall_assessment": "string",
                "compliance_observations": [],
                "risk_level": "low|medium|high",
                "risks": [],
                "missing_provisions": [],
                "one_sided_terms": [],
                "key_legal_terms": [],
                "referenced_laws": [],
                "clauses": [
                    {
                        "title": "string",
                        "summary": "string",
                        "risk_notes": "string",
                    }
                ],
                "recommendations": [],
            },
            "document_text": (text or "")[:12000],
            "response_rule": "Return JSON only.",
        }
    )

    try:
        response = requests.post(
            f"{getattr(settings, 'OLLAMA_URL', 'http://localhost:11434')}/api/generate",
            json={
                "model": getattr(settings, "OLLAMA_MODEL", "llama3"),
                "prompt": prompt,
                "stream": False,
                "format": "json",
            },
            timeout=180,
        )
        response.raise_for_status()
        parsed = json.loads(response.json().get("response", "{}"))
    except (requests.RequestException, ValueError, json.JSONDecodeError):
        parsed = {}

    return {
        "detected_document_type": parsed.get("detected_document_type", document_type),
        "executive_summary": parsed.get("executive_summary", ""),
        "overall_assessment": parsed.get("overall_assessment", ""),
        "compliance_observations": parsed.get("compliance_observations", []),
        "risk_level": parsed.get("risk_level", "medium"),
        "risks": parsed.get("risks", []),
        "missing_provisions": parsed.get("missing_provisions", []),
        "one_sided_terms": parsed.get("one_sided_terms", []),
        "key_legal_terms": parsed.get("key_legal_terms", []),
        "referenced_laws": parsed.get("referenced_laws", []),
        "clauses": parsed.get("clauses", []),
        "recommendations": parsed.get("recommendations", []),
    }
